"""Генерирует спецификации разъёмов и пинов из tools/data.py.

Запуск из корня репозитория:  python3 tools/build_specs.py
Результат: specs/Wiring_N/Wiring_N_spec.xlsx и .md, specs/Summary_all_drawings.xlsx,
drawings_registry.csv, README.md.
"""
import csv
import os
import sys
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from data import CATALOG, COMMON_REMARKS, DRAWINGS, STATUS, TERM_RANGE  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPECS = os.path.join(ROOT, "specs")

HDR = PatternFill("solid", fgColor="1F4E78")
FILL = {"Подтверждён": PatternFill("solid", fgColor="E2EFDA"), "Вероятно": PatternFill("solid", fgColor="FFF2CC"),
        "Проверить": PatternFill("solid", fgColor="F8CBAD")}
WARN = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ICON = {"Подтверждён": "✅", "Вероятно": "⚠️", "Проверить": "❌", "—": ""}


def car(d):
    return f"{d['brand']} {d['model']}" if d["brand"] != "?" else "модель не указана"


def status(pn):
    return STATUS[CATALOG[pn][4]] if pn in CATALOG else "—"


def fmt_mm2(v):
    return "" if v in ("", None) else f"{v:g}"


def check_sheet(sh):
    """Возвращает {(ref, cav): (двойной обжим, проблема сечения)} и список текстовых замечаний."""
    conns = {c["ref"]: c for c in sh["connectors"]}
    groups = OrderedDict()
    for p in sh["pins"]:
        if p["mm2"] == "":
            continue
        c = conns[p["ref"]]
        if p["cav"] == "—" and not (c["cavities"] == 1 and c["term"]):
            key = (p["ref"], "—", p["wire"])  # отдельный наконечник на каждый провод
        else:
            key = (p["ref"], p["cav"])
        groups.setdefault(key, []).append(p)
    flags, notes = {}, []
    for key, rows in groups.items():
        term = conns[key[0]]["term"]
        total = round(sum(r["mm2"] for r in rows), 2)
        double = len(rows) > 1
        problem = ""
        if term in TERM_RANGE:
            lo, hi = TERM_RANGE[term]
            if total > hi:
                problem = f"{total:g} мм² > {hi:g} мм² (макс. {term})"
            elif total < lo:
                problem = f"{total:g} мм² < {lo:g} мм² (мин. {term})"
        flags[key] = (double, problem)
        where = f"{key[0]}" + (f", пин {key[1]}" if key[1] != "—" else "")
        if double:
            wires = " + ".join(f"{r['wire']} ({r['color'].lower()} {r['mm2']:g})" for r in rows)
            notes.append(f"Лист {sh['sheet']}, {where}: двойной обжим {wires} = {total:g} мм²"
                         + (f" — {problem}" if problem else ""))
        elif problem:
            notes.append(f"Лист {sh['sheet']}, {where}: провод {rows[0]['wire']} — {problem}")
    return flags, notes


def pin_flags(sh, flags, p):
    c = {c["ref"]: c for c in sh["connectors"]}[p["ref"]]
    if p["mm2"] == "":
        return "", ""
    if p["cav"] == "—" and not (c["cavities"] == 1 and c["term"]):
        key = (p["ref"], "—", p["wire"])
    else:
        key = (p["ref"], p["cav"])
    d, prob = flags.get(key, (False, ""))
    return ("Да" if d else ""), prob


def collect(d):
    """Плоские таблицы по одному чертежу."""
    conn_rows, pin_rows, bom_rows, remarks, used = [], [], [], [], OrderedDict()
    for sh in d["sheets"]:
        flags, notes = check_sheet(sh)
        for c in sh["connectors"]:
            pn = c["pn"] or "—"
            term = c["term"] or "—"
            st = status(c["pn"]) if c["pn"] else status(c["term"]) if c["term"] else "—"
            conn_rows.append([d["no"], car(d), sh["sheet"], c["ref"], c["pos"], c["kind"], pn,
                              CATALOG.get(c["pn"], ("—",))[0] if c["pn"] else "—", c["cavities"], term,
                              c["term_qty"] or "—", st, c["note"]])
            for x in (c["pn"], c["term"]):
                if x in CATALOG:
                    used[x] = True
        for p in sh["pins"]:
            c = {c["ref"]: c for c in sh["connectors"]}[p["ref"]]
            dbl, prob = pin_flags(sh, flags, p)
            pin_rows.append([d["no"], car(d), sh["sheet"], p["ref"], c["pn"] or "—", c["term"] or "—", p["cav"], p["wire"],
                             p["color"], fmt_mm2(p["mm2"]), p["to"], dbl, prob, p["note"]])
        for pos, name, pn, qty, unit in sh["bom"]:
            bom_rows.append([d["no"], car(d), sh["sheet"], pos, name, pn, CATALOG.get(pn, ("—",))[0], qty, unit, status(pn)])
            if pn in CATALOG:
                used[pn] = True
        remarks += [("Проверка сечений", n) for n in notes]
        remarks += [(f"Лист {sh['sheet']}", r) for r in sh["remarks"]]
    return conn_rows, pin_rows, bom_rows, remarks, list(used)


CONN_HDR = ["№ чертежа", "Автомобиль", "Лист", "Разъём", "Поз. BOM", "Назначение", "P/N корпуса", "Производитель",
            "Позиций", "Контакт (P/N)", "Контактов, шт", "Статус", "Примечание"]
CONN_W = [9, 22, 9, 18, 7, 40, 22, 20, 9, 22, 10, 13, 50]
PIN_HDR = ["№ чертежа", "Автомобиль", "Лист", "Разъём", "P/N корпуса", "Контакт", "Пин", "Провод №", "Цвет",
           "Сечение, мм²", "Куда идёт", "Двойной обжим", "Проверка сечения", "Примечание"]
PIN_W = [9, 22, 9, 18, 20, 22, 6, 9, 12, 9, 26, 9, 34, 50]
BOM_HDR = ["№ чертежа", "Автомобиль", "Лист", "Поз.", "Наименование", "P/N", "Производитель", "Кол-во", "Ед.", "Статус"]
BOM_W = [9, 22, 9, 6, 36, 24, 20, 8, 6, 13]
CAT_HDR = ["P/N", "Производитель", "Серия", "Описание", "Сечение, мм²", "Статус", "Аналог / комментарий", "Источник"]
CAT_W = [24, 22, 30, 50, 16, 13, 55, 60]


def write_sheet(wb, title, header, rows, widths, status_col=None, first=False, hilite=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.border = BORDER
            if c.row > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        if hilite:
            fill = hilite(row)
            if fill:
                for c in row:
                    c.fill = fill
        if status_col is not None and row[status_col].value in FILL:
            row[status_col].fill = FILL[row[status_col].value]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def pin_hilite(row):
    if row[12].value:
        return BAD
    if row[11].value == "Да":
        return WARN
    return None


def cat_rows(pns):
    return [[pn, *CATALOG[pn][:4], STATUS[CATALOG[pn][4]], *CATALOG[pn][5:]] for pn in pns]


def common_for(no):
    out = []
    for title, where, text, rec in COMMON_REMARKS:
        nums = set()
        for part in where.split(","):
            part = part.strip()
            if "–" in part:
                a, b = part.split("–")
                nums |= set(range(int(a), int(b) + 1))
            else:
                nums.add(int(part))
        if no in nums:
            out.append((title, text, rec))
    return out


def build_drawing(d):
    conn_rows, pin_rows, bom_rows, remarks, used = collect(d)
    folder = os.path.join(SPECS, d["folder"])
    os.makedirs(folder, exist_ok=True)
    wb = Workbook()
    info = [["№ чертежа", d["no"]], ["Марка", d["brand"]], ["Модель", d["model"]]]
    info += [[f"Лист {sh['sheet']}", f"{sh['desc']}  (файл {d['sources'][sh['sheet']]})"] for sh in d["sheets"]]
    write_sheet(wb, "Чертёж", ["Параметр", "Значение"], info, [18, 110], first=True)
    write_sheet(wb, "Разъёмы", CONN_HDR, conn_rows, CONN_W, status_col=11)
    write_sheet(wb, "Распиновка", PIN_HDR, pin_rows, PIN_W, hilite=pin_hilite)
    write_sheet(wb, "BOM", BOM_HDR, bom_rows, BOM_W, status_col=9)
    rem = [[i, t, x, ""] for i, (t, x) in enumerate(remarks, 1)]
    rem += [[len(rem) + i, t, x, r] for i, (t, x, r) in enumerate(common_for(d["no"]), 1)]
    write_sheet(wb, "Замечания", ["№", "Тема", "Замечание", "Рекомендация"], rem, [5, 28, 90, 60])
    write_sheet(wb, "Каталог", CAT_HDR, cat_rows(used), CAT_W, status_col=5)
    wb.save(os.path.join(folder, f"{d['folder']}_spec.xlsx"))
    write_md(d, conn_rows, pin_rows, bom_rows, rem, used, folder)
    return conn_rows, pin_rows, bom_rows, rem, used


def md_table(header, rows):
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for r in rows:
        out.append("| " + " | ".join(str(x).replace("|", "/") for x in r) + " |")
    return "\n".join(out)


def write_md(d, conn_rows, pin_rows, bom_rows, rem, used, folder):
    L = [f"# {d['folder']} — спецификация разъёмов и пинов", "",
         md_table(["№ чертежа", "Марка", "Модель"], [[d["no"], d["brand"], d["model"]]]), "",
         "Статусы: ✅ подтверждён по каталогу · ⚠️ вероятно · ❌ проверить. "
         "В распиновке: **ДО** — двойной обжим, 🔴 — сечение вне диапазона контакта.", ""]
    for sh in d["sheets"]:
        s = sh["sheet"]
        L += [f"## Лист {s}", "", f"{sh['desc']}. Исходник: `{d['sources'][s]}`", "", "**Разъёмы**", ""]
        L.append(md_table(["Разъём", "Поз.", "Назначение", "P/N корпуса", "Поз-й", "Контакт", "Шт", "Статус"],
                          [[r[3], r[4], r[5], r[6], r[8], r[9], r[10], ICON.get(r[11], "") + " " + r[11]]
                           for r in conn_rows if r[2] == s]))
        L += ["", "**Распиновка**", ""]
        L.append(md_table(["Разъём", "Пин", "Провод", "Цвет", "мм²", "Куда", "Флаги", "Примечание"],
                          [[r[3], r[6], r[7], r[8], r[9], r[10],
                            ("ДО " if r[11] else "") + ("🔴 " + r[12] if r[12] else ""), r[13]]
                           for r in pin_rows if r[2] == s]))
        L += ["", "**BOM листа**", ""]
        L.append(md_table(["Поз.", "Наименование", "P/N", "Кол-во", "Статус"],
                          [[r[3], r[4], r[5], f"{r[7]:g} {r[8]}", ICON.get(r[9], "")] for r in bom_rows if r[2] == s]))
        L.append("")
    L += ["## Замечания", ""]
    L += [f"{r[0]}. **{r[1]}.** {r[2]}" + (f" → {r[3]}" if r[3] else "") for r in rem]
    L += ["", "## Каталог использованных номеров", ""]
    L.append(md_table(["P/N", "Производитель", "Описание", "мм²", "Статус", "Комментарий"],
                      [[pn, CATALOG[pn][0], CATALOG[pn][2], CATALOG[pn][3], ICON[STATUS[CATALOG[pn][4]]] if CATALOG[pn][4] != "-" else "",
                        CATALOG[pn][5]] for pn in used]))
    srcs = [(pn, CATALOG[pn][6]) for pn in used if CATALOG[pn][6]]
    if srcs:
        L += ["", "## Источники", ""] + [f"- [{pn}]({u})" for pn, u in srcs]
    with open(os.path.join(folder, f"{d['folder']}_spec.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")


def build_summary(results):
    wb = Workbook()
    reg = [[d["no"], d["brand"], d["model"], len(d["sheets"]), f"specs/{d['folder']}"] for d in DRAWINGS]
    write_sheet(wb, "Реестр", ["№ чертежа", "Марка", "Модель", "Листов", "Папка"], reg, [10, 10, 34, 8, 20], first=True)
    # потребность по номерам: P/N x чертёж (1 комплект каждого листа)
    qty = defaultdict(lambda: defaultdict(float))
    for d, (_, _, bom, _, _) in zip(DRAWINGS, results):
        for r in bom:
            if r[8] == "шт":
                qty[r[5]][d["no"]] += r[7]
    hdr = ["P/N", "Производитель", "Описание", "Статус"] + [f"№{d['no']} {d['model']}" for d in DRAWINGS] + ["Итого"]
    rows = []
    for pn in sorted(qty, key=lambda x: (CATALOG.get(x, ("",))[0], x)):
        per = [qty[pn].get(d["no"], "") for d in DRAWINGS]
        cat = CATALOG.get(pn, ("—", "", "—", "", "-"))
        rows.append([pn, cat[0], cat[2], STATUS[cat[4]], *[int(x) if x != "" else "" for x in per],
                     int(sum(x for x in per if x != ""))])
    write_sheet(wb, "Потребность по моделям", hdr, rows, [24, 22, 46, 13] + [14] * len(DRAWINGS) + [8], status_col=3)
    write_sheet(wb, "Каталог", CAT_HDR, cat_rows([pn for pn in CATALOG]), CAT_W, status_col=5)
    write_sheet(wb, "Все разъёмы", CONN_HDR, [r for res in results for r in res[0]], CONN_W, status_col=11)
    write_sheet(wb, "Вся распиновка", PIN_HDR, [r for res in results for r in res[1]], PIN_W, hilite=pin_hilite)
    rem = []
    for d, res in zip(DRAWINGS, results):
        rem += [[d["no"], car(d), r[1], r[2], r[3]] for r in res[3]]
    write_sheet(wb, "Все замечания", ["№ чертежа", "Автомобиль", "Тема", "Замечание", "Рекомендация"], rem,
                [9, 22, 28, 90, 60])
    wb.save(os.path.join(SPECS, "Summary_all_drawings.xlsx"))
    return qty


def write_registry():
    with open(os.path.join(ROOT, "drawings_registry.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["drawing_no", "brand", "model", "sheets", "folder", "status"])
        for d in DRAWINGS:
            w.writerow([d["no"], d["brand"], d["model"], ";".join(s["sheet"] for s in d["sheets"]),
                        f"specs/{d['folder']}", "done"])
    rows = [[d["no"], d["brand"], d["model"], len(d["sheets"]),
             f"[{d['folder']}](specs/{d['folder']}/{d['folder']}_spec.md)"] for d in DRAWINGS]
    readme = f"""# Model_Electric

Спецификации разъёмов и пинов по чертежам жгутов. Каждый чертёж привязан к модели автомобиля.

## Реестр чертежей

{md_table(["№ чертежа", "Марка", "Модель", "Листов", "Спецификация"], rows)}

Машиночитаемая версия реестра: [`drawings_registry.csv`](drawings_registry.csv).
Сводка по всем чертежам (потребность по моделям, каталог, вся распиновка, все замечания):
[`specs/Summary_all_drawings.xlsx`](specs/Summary_all_drawings.xlsx).

## Структура

`specs/Wiring_N/`:
- `Wiring_N[-k]_source.pdf` — исходные листы чертежа;
- `Wiring_N_spec.xlsx` — листы «Чертёж», «Разъёмы», «Распиновка», «BOM», «Замечания», «Каталог». В каждой строке есть «№ чертежа» и «Автомобиль»;
- `Wiring_N_spec.md` — то же в markdown.

`tools/data.py` — данные, считанные с чертежей, и каталог номеров. `tools/build_specs.py` — генератор:

```bash
python3 tools/build_specs.py
```
"""
    with open(os.path.join(ROOT, "README.md"), "w", encoding="utf-8") as f:
        f.write(readme)


def main():
    results = [build_drawing(d) for d in DRAWINGS]
    build_summary(results)
    write_registry()
    for d, r in zip(DRAWINGS, results):
        print(f"№{d['no']} {car(d)}: разъёмов {len(r[0])}, пинов {len(r[1])}, замечаний {len(r[3])}")


if __name__ == "__main__":
    main()
